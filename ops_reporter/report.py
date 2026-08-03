from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .core import calculate_metrics, daily_summary, department_summary

HEADER_FILL = PatternFill("solid", fgColor="17365D")
ACCENT_FILL = PatternFill("solid", fgColor="D9EAF7")
WHITE_FONT = Font(color="FFFFFF", bold=True)


def _style_table(sheet) -> None:
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = WHITE_FONT
        cell.alignment = Alignment(horizontal="center")
    sheet.freeze_panes = "A2"
    for column in sheet.columns:
        width = min(max(len(str(cell.value or "")) for cell in column) + 2, 34)
        sheet.column_dimensions[get_column_letter(column[0].column)].width = width


def generate_excel_report(frame: pd.DataFrame, output_path: str | Path) -> Path:
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    metrics = calculate_metrics(frame)
    daily = daily_summary(frame)
    departments = department_summary(frame)

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        daily.to_excel(writer, sheet_name="Daily Metrics", index=False)
        departments.to_excel(writer, sheet_name="Department Summary", index=False)
        frame.to_excel(writer, sheet_name="Source Data", index=False)
        dashboard = writer.book.create_sheet("Dashboard", 0)
        dashboard.append(["Operations Performance Report"])
        dashboard.merge_cells("A1:D1")
        dashboard["A1"].fill = HEADER_FILL
        dashboard["A1"].font = Font(color="FFFFFF", bold=True, size=18)
        dashboard["A1"].alignment = Alignment(horizontal="center")
        dashboard.append([])
        dashboard.append(["KPI", "Value", "KPI", "Value"])
        dashboard.append(["Total units", metrics.total_units, "Target units", metrics.target_units])
        dashboard.append(["Quality rate", metrics.quality_rate, "Target attainment", metrics.target_attainment])
        dashboard.append(["Total defects", metrics.total_defects, "Downtime (minutes)", metrics.total_downtime_minutes])
        for cell in dashboard[3]:
            cell.fill = ACCENT_FILL
            cell.font = Font(bold=True)
        dashboard["B5"].number_format = "0.0%"
        dashboard["D5"].number_format = "0.0%"
        for col in "ABCD":
            dashboard.column_dimensions[col].width = 22

        for sheet_name in ["Daily Metrics", "Department Summary", "Source Data"]:
            _style_table(writer.book[sheet_name])
        for row in range(2, writer.book["Daily Metrics"].max_row + 1):
            writer.book["Daily Metrics"].cell(row, 1).number_format = "yyyy-mm-dd"
        for row in range(2, writer.book["Source Data"].max_row + 1):
            writer.book["Source Data"].cell(row, 1).number_format = "yyyy-mm-dd"
        for sheet_name in ["Daily Metrics", "Department Summary"]:
            sheet = writer.book[sheet_name]
            for row in range(2, sheet.max_row + 1):
                sheet.cell(row, 6).number_format = "0.0%"
                sheet.cell(row, 7).number_format = "0.0%"
            sheet.conditional_formatting.add(
                f"G2:G{sheet.max_row}",
                ColorScaleRule(start_type="min", start_color="F8696B", mid_type="percentile", mid_value=50, mid_color="FFEB84", end_type="max", end_color="63BE7B"),
            )

        daily_sheet = writer.book["Daily Metrics"]
        line = LineChart()
        line.title = "Production vs target"
        line.y_axis.title = "Units"
        line.x_axis.title = "Date"
        line.add_data(Reference(daily_sheet, min_col=2, max_col=2, min_row=1, max_row=daily_sheet.max_row), titles_from_data=True)
        line.add_data(Reference(daily_sheet, min_col=5, max_col=5, min_row=1, max_row=daily_sheet.max_row), titles_from_data=True)
        line.set_categories(Reference(daily_sheet, min_col=1, min_row=2, max_row=daily_sheet.max_row))
        line.x_axis.number_format = "yyyy-mm-dd"
        dashboard.add_chart(line, "A9")

        dept_sheet = writer.book["Department Summary"]
        bar = BarChart()
        bar.title = "Downtime by department"
        bar.y_axis.title = "Minutes"
        bar.add_data(Reference(dept_sheet, min_col=4, min_row=1, max_row=dept_sheet.max_row), titles_from_data=True)
        bar.set_categories(Reference(dept_sheet, min_col=1, min_row=2, max_row=dept_sheet.max_row))
        dashboard.add_chart(bar, "H9")

    return output


def generate_html_summary(frame: pd.DataFrame, output_path: str | Path) -> Path:
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    metrics = calculate_metrics(frame)
    departments = department_summary(frame)
    html = f"""<!doctype html><html><head><meta charset='utf-8'><title>Operations Report</title><style>body{{font-family:Arial;max-width:1100px;margin:40px auto;color:#172033}}.kpis{{display:grid;grid-template-columns:repeat(3,1fr);gap:14px}}.kpi{{padding:18px;border:1px solid #ccd6e2;border-radius:12px}}table{{border-collapse:collapse;width:100%;margin-top:28px}}th,td{{padding:10px;border-bottom:1px solid #dbe3ec;text-align:right}}th:first-child,td:first-child{{text-align:left}}</style></head><body><h1>Operations Performance Report</h1><div class='kpis'><div class='kpi'><b>Total units</b><br>{metrics.total_units:,}</div><div class='kpi'><b>Quality rate</b><br>{metrics.quality_rate:.1%}</div><div class='kpi'><b>Target attainment</b><br>{metrics.target_attainment:.1%}</div><div class='kpi'><b>Defects</b><br>{metrics.total_defects:,}</div><div class='kpi'><b>Downtime</b><br>{metrics.total_downtime_minutes:,} min</div><div class='kpi'><b>Target units</b><br>{metrics.target_units:,}</div></div>{departments.to_html(index=False, float_format=lambda value: f'{value:.1%}' if value <= 2 else f'{value:,.0f}')} </body></html>"""
    output.write_text(html, encoding="utf-8")
    return output
