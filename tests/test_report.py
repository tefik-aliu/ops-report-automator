from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from ops_reporter.core import validate_data
from ops_reporter.report import generate_excel_report, generate_html_summary


def test_reports_are_created(tmp_path: Path):
    frame = validate_data(pd.DataFrame([
        {"date":"2026-07-01","department":"Cutting","operator":"A","units_produced":100,"defects":2,"downtime_minutes":10,"target_units":110},
        {"date":"2026-07-02","department":"Assembly","operator":"B","units_produced":80,"defects":1,"downtime_minutes":5,"target_units":80},
    ]))
    excel = generate_excel_report(frame, tmp_path / "report.xlsx")
    html = generate_html_summary(frame, tmp_path / "report.html")
    assert excel.exists() and html.exists()
    workbook = load_workbook(excel, read_only=True)
    assert workbook.sheetnames == ["Dashboard", "Daily Metrics", "Department Summary", "Source Data"]
    assert "Operations Performance Report" in html.read_text(encoding="utf-8")
